import streamlit as st
import io
import re
import time
import pandas as pd
import pdfplumber
from datetime import datetime
from openpyxl.styles import Border, Side, Alignment, Font

# --- HELPER: TIME FORMATTER ---
def format_seconds(sec):
    if sec < 60:
        return f"{sec:.1f} sec"
    mins = int(sec // 60)
    rem_sec = sec % 60
    return f"{mins}m {rem_sec:.1f}s"

# --- HELPER: UNIVERSAL DATE PARSER ---
def parse_any_date(date_str):
    clean_d = re.sub(r'\s+', '-', date_str.strip()).replace('/', '-')
    for fmt in ('%d-%m-%Y', '%d-%m-%y', '%d-%b-%Y', '%d-%b-%y', '%d-%B-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(clean_d, fmt)
        except ValueError:
            continue
    return None

# --- UNIVERSAL METADATA EXTRACTOR ---
def extract_metadata(text):
    flat_text = text.replace('\n', '  ')
    meta = {
        "name": "-", "address": "-", "cust_id": "-", "ifsc": "-", 
        "acct_num": "-", "branch": "-", "date_from": "-", "date_to": "-"
    }
    
    name_m = re.search(r'(?:Account\s*holder\s*name|Customer\s*Name|Account\s*Name|Name)\s*[:\-]\s*(.*?)\s+(?:Account|Customer|Address|Branch|IFSC|CIF)', flat_text, re.IGNORECASE)
    if name_m: meta["name"] = name_m.group(1).strip()
    
    addr_m = re.search(r'(?:Account\s*holder\s*address|Customer\s*Address|Address)\s*[:\-]\s*(.*?)\s+(?:Customer|IFSC|Account|Branch|Email|Mobile|Phone)', flat_text, re.IGNORECASE)
    if addr_m: meta["address"] = addr_m.group(1).strip()
    
    cust_m = re.search(r'(?:Customer\s*ID|Cust\s*ID|CIF\s*No|CIF\s*ID)\s*[:\-]?\s*([A-Za-z0-9]+)', flat_text, re.IGNORECASE)
    if cust_m: meta["cust_id"] = cust_m.group(1).strip()
    
    ifsc_m = re.search(r'(?:IFSC\s*Code|RTGS\/NEFT\s*IFSC|IFSC)\s*[:\-]?\s*([A-Z]{4}0[A-Z0-9]{6})', flat_text, re.IGNORECASE)
    if ifsc_m: meta["ifsc"] = ifsc_m.group(1).strip()
    
    acct_m = re.search(r'(?:Account\s*number|Account\s*No|A\/c\s*No|Ac\s*No)\s*[:\-]?\s*(\d{9,18})', flat_text, re.IGNORECASE)
    if acct_m: meta["acct_num"] = acct_m.group(1).strip()
    
    branch_m = re.search(r'(?:Branch\s*Name|Home\s*Branch|Branch)\s*[:\-]\s*([A-Za-z0-9\s,\-]+?)(?:\s{2,}|IFSC|Account|Customer|Date|Period)', flat_text, re.IGNORECASE)
    if branch_m: meta["branch"] = branch_m.group(1).strip()
    
    from_m = re.search(r'(?:from|Period)\s*[:\-]?\s*(\d{1,2}[-/\s][A-Za-z0-9]{2,3}[-/\s]\d{2,4})\s+to\s*[:\-]?\s*(\d{1,2}[-/\s][A-Za-z0-9]{2,3}[-/\s]\d{2,4})', flat_text, re.IGNORECASE)
    if from_m: 
        meta["date_from"] = from_m.group(1).strip()
        meta["date_to"] = from_m.group(2).strip()
        
    return meta

# --- UNIVERSAL MULTI-BANK CONVERSION ENGINE ---
def process_pdf(pdf_file, pdf_password, progress_bar, status_text, timer_text):
    start_time = time.time()
    full_text = ""
    
    # 1. Open PDF
    try:
        pdf = pdfplumber.open(pdf_file, password=pdf_password if pdf_password else None)
        if len(pdf.pages) > 0:
            _ = pdf.pages[0].chars # Trigger decryption check
    except pdfplumber.pdfminer.pdfdocument.PDFPasswordIncorrect:
        raise ValueError("Incorrect password or PDF is password-protected. Please enter the correct password.")
    except Exception as e:
        raise ValueError(f"Failed to open PDF: {e}")

    # 2. Extract Text
    total_pages = len(pdf.pages)
    for idx, page in enumerate(pdf.pages):
        text = page.extract_text()
        if text:
            full_text += text + "\n"
            
        elapsed = time.time() - start_time
        pct = int(((idx + 1) / total_pages) * 75)
        
        progress_bar.progress(pct)
        status_text.text(f"Scanning PDF: Page {idx + 1} of {total_pages} ({pct}%)")
        timer_text.text(f"Time Taken: {format_seconds(elapsed)}")
    
    pdf.close()

    # 3. Parse Transactions
    status_text.text("Validating transactions & calculating balances (82%)...")
    progress_bar.progress(82)

    lines = full_text.split('\n')
    meta = extract_metadata(full_text)

    transactions = []
    current_txn = None

    date_rgx = r'(?:\d{1,2}[-/]\d{1,2}[-/]\d{2,4}|\d{1,2}[-\s][A-Za-z]{3}[-\s]\d{2,4})'
    line_start_pattern = re.compile(rf'^(?:\d+\s+)?({date_rgx})(?:\s+{date_rgx})?\s+(.+)$')
    
    amt_rgx = r'(?:₹|Rs\.?|INR)?\s*([\d,]+\.\d{2})\s*(?:Cr|Dr|CR|DR|cr|dr)?'
    three_amts_pattern = re.compile(rf'^(.*?)\s+{amt_rgx}\s+{amt_rgx}\s+{amt_rgx}$')
    two_amts_pattern = re.compile(rf'^(.*?)\s+{amt_rgx}\s+{amt_rgx}$')
    one_amt_pattern = re.compile(rf'^(.*?)\s+{amt_rgx}$')
    standalone_bal_pattern = re.compile(rf'^{amt_rgx}$')
    wrapped_two_amts_pattern = re.compile(rf'^(.*?)\s*{amt_rgx}\s+{amt_rgx}$')

    ignore_words = [
        "discrepancy", "passwords", "atm", "page ", "statement of", "opening balance",
        "closing balance", "total", "generated on", "computer generated", "branch",
        "date remarks", "date narration", "date particulars", "withdrawal", "deposit",
        "chq/ref", "cheque", "brought forward", "carried forward"
    ]

    for line in lines:
        line = line.strip()
        if not line or line.startswith('---'):
            continue

        start_match = line_start_pattern.match(line)
        if start_match:
            if current_txn and current_txn['Balance'] is not None and current_txn['Amount'] is not None:
                transactions.append(current_txn)

            txn_date = start_match.group(1)
            rest = start_match.group(2).strip()

            m3 = three_amts_pattern.match(rest)
            m2 = two_amts_pattern.match(rest)
            m1 = one_amt_pattern.match(rest)

            if m3:
                current_txn = {'Date': txn_date, 'Remarks': m3.group(1).strip(), 
                               'Amount': float(m3.group(2).replace(',', '')) if float(m3.group(2).replace(',', '')) > 0 else float(m3.group(3).replace(',', '')), 
                               'Balance': float(m3.group(4).replace(',', ''))}
            elif m2:
                current_txn = {'Date': txn_date, 'Remarks': m2.group(1).strip(), 
                               'Amount': float(m2.group(2).replace(',', '')), 'Balance': float(m2.group(3).replace(',', ''))}
            elif m1:
                current_txn = {'Date': txn_date, 'Remarks': m1.group(1).strip(), 
                               'Amount': float(m1.group(2).replace(',', '')), 'Balance': None}
            else:
                current_txn = {'Date': txn_date, 'Remarks': rest, 'Amount': None, 'Balance': None}
            continue

        if current_txn and current_txn['Balance'] is None:
            bal_match = standalone_bal_pattern.match(line)
            if bal_match and current_txn['Amount'] is not None:
                current_txn['Balance'] = float(bal_match.group(1).replace(',', ''))
                continue
                
            wrap_m2 = wrapped_two_amts_pattern.match(line)
            if wrap_m2 and current_txn['Amount'] is None:
                extra_rem = wrap_m2.group(1).strip()
                if extra_rem: current_txn['Remarks'] += " " + extra_rem
                current_txn['Amount'] = float(wrap_m2.group(2).replace(',', ''))
                current_txn['Balance'] = float(wrap_m2.group(3).replace(',', ''))
                continue

        if current_txn:
            lower_line = line.lower()
            if any(w in lower_line for w in ignore_words):
                continue
            current_txn['Remarks'] = (current_txn['Remarks'] + " " + line).strip()

    if current_txn and current_txn['Balance'] is not None and current_txn['Amount'] is not None:
        transactions.append(current_txn)

    if not transactions:
        raise ValueError("Could not detect transactions. Ensure it is a text-based bank statement.")

    # 4. Chronological Check
    if len(transactions) > 1:
        first_dt = parse_any_date(transactions[0]['Date'])
        last_dt = parse_any_date(transactions[-1]['Date'])
        if first_dt and last_dt and first_dt > last_dt:
            transactions.reverse()

    df = pd.DataFrame(transactions)

    # 5. Math Validation
    debits, credits = [], []
    for i in range(len(df)):
        amt = df.loc[i, 'Amount']
        remarks = df.loc[i, 'Remarks'].upper()
        if i > 0:
            diff = round(df.loc[i, 'Balance'] - df.loc[i-1, 'Balance'], 2)
            if diff > 0:
                credits.append(amt); debits.append(None)
            elif diff < 0:
                debits.append(amt); credits.append(None)
            else:
                if '/CR/' in remarks or 'CREDIT' in remarks or 'BY ' in remarks or 'NEFT CR' in remarks:
                    credits.append(amt); debits.append(None)
                else:
                    debits.append(amt); credits.append(None)
        else:
            if '/CR/' in remarks or 'CREDIT' in remarks or 'BY ' in remarks or 'NEFT CR' in remarks:
                credits.append(amt); debits.append(None)
            else:
                debits.append(amt); credits.append(None)

    df['Debit'] = debits; df['Credit'] = credits
    df = df[['Date', 'Remarks', 'Debit', 'Credit', 'Balance']]

    # 6. Build Excel File to Memory Buffer
    status_text.text("Applying professional formatting (92%)...")
    progress_bar.progress(92)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Statement', index=False, startrow=13)
        worksheet = writer.sheets['Statement']

        worksheet.merge_cells('A1:E2')
        worksheet['A1'] = "Detailed Statement"
        worksheet['A1'].font = Font(size=18, bold=True)
        worksheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
        worksheet['E3'] = f"Date: {datetime.today().strftime('%d/%m/%Y')}"
        worksheet['E3'].font = Font(bold=True)
        worksheet['E3'].alignment = Alignment(horizontal='right')

        worksheet['A4'] = "Account holder name:"; worksheet['B4'] = meta["name"]
        worksheet['C4'] = "Account holder address:"; worksheet.merge_cells('D4:E4'); worksheet['D4'] = meta["address"]
        worksheet['A5'] = "Customer ID:"; worksheet['B5'] = meta["cust_id"]
        worksheet['C5'] = "IFSC:"; worksheet.merge_cells('D5:E5'); worksheet['D5'] = meta["ifsc"]
        worksheet['A6'] = "Account number:"; worksheet['B6'] = meta["acct_num"]
        worksheet['C6'] = "Branch Name:"; worksheet.merge_cells('D6:E6'); worksheet['D6'] = meta["branch"]
        
        box_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        for r in range(4, 7):
            for c in range(1, 6):
                cell = worksheet.cell(row=r, column=c)
                cell.border = box_border
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if c in [1, 3]: cell.font = Font(size=10)
                else: cell.font = Font(size=10, bold=True)

        worksheet['A8'] = "Transaction Date"; worksheet['B8'] = f"from: {meta['date_from']}"; worksheet['C8'] = f"to: {meta['date_to']}"
        worksheet['A9'] = "Amount"; worksheet['B9'] = "from: -"; worksheet['C9'] = "to: -"
        worksheet['A10'] = "Cheque"; worksheet['B10'] = "from: -"; worksheet['C10'] = "to: -"
        worksheet['A11'] = "Transaction type: All"
        for r in range(8, 12):
            for c in range(1, 4): worksheet.cell(row=r, column=c).font = Font(bold=True, size=10)

        headers = ["DATE", "REMARKS", "DEBIT", "CREDIT", "BALANCE"]
        for col_num, header in enumerate(headers, 1):
            cell = worksheet.cell(row=14, column=col_num, value=header)
            cell.font = Font(bold=True); cell.border = box_border; cell.alignment = Alignment(horizontal='center', vertical='center')

        max_row = worksheet.max_row
        for row in worksheet.iter_rows(min_row=15, max_row=max_row, min_col=1, max_col=5):
            for cell in row:
                cell.border = box_border
                cell.alignment = Alignment(wrap_text=True, vertical='center')

        for r in range(16, max_row + 1):
            worksheet[f"E{r}"] = f"=E{r-1}+SUM(D{r})-SUM(C{r})"

        worksheet.column_dimensions['A'].width = 18
        worksheet.column_dimensions['B'].width = 45
        worksheet.column_dimensions['C'].width = 15
        worksheet.column_dimensions['D'].width = 15
        worksheet.column_dimensions['E'].width = 18
        worksheet.row_dimensions[4].height = 30

    progress_bar.progress(100)
    status_text.text(f"✅ Complete! Extracted {len(df)} transactions.")
    output.seek(0)
    return output

# --- STREAMLIT UI ---
st.set_page_config(page_title="Universal Bank Converter", page_icon="🏦")
st.title("🏦 Universal Bank Statement Converter")
st.write("Convert your PDF bank statements directly to a formatted Excel file.")

uploaded_file = st.file_uploader("Upload Bank Statement PDF", type=["pdf"])
pdf_password = st.text_input("PDF Password (leave blank if none)", type="password")

if uploaded_file is not None:
    if st.button("Convert to Excel"):
        progress_bar = st.progress(0)
        status_text = st.empty()
        timer_text = st.empty()
        
        try:
            excel_buffer = process_pdf(uploaded_file, pdf_password, progress_bar, status_text, timer_text)
            
            st.success("Conversion successful!")
            st.download_button(
                label="📥 Download Excel File",
                data=excel_buffer,
                file_name=f"{uploaded_file.name.replace('.pdf', '')}_Converted.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            st.error(f"Error: {e}")
